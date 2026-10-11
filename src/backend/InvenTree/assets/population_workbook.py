"""Parse a UI population workbook into validated, Django-free records.

The workbook is the Cedar Creek "AIMMS UI population master" layout: one
sheet per staging table, first row is the header. Only the sheets a
population run needs are read; everything else in the file is ignored.

Nothing here touches the database. The parser's job is to turn a
spreadsheet into typed records and to refuse anything that could not be
applied safely: duplicate codes, dangling parents, a station with no
identity, a pump slot whose key the registry would reject. Catching those
here means the planner can assume a coherent input and spend its checks
on the database side instead.
"""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from uuid import UUID

MAX_WORKBOOK_BYTES = 64 * 1024 * 1024
MAX_ROWS_PER_SHEET = 10_000

SHEET_LOCATIONS = 'Facility Locations'
SHEET_MACHINES = 'Machine Records'
SHEET_PROFILES = 'Machine Profiles'
SHEET_COMPONENTS = 'Component Occurrences'
SHEET_PART_LINKS = 'Machine Part Links'
SHEET_PARTS = 'Parts Catalogue'
SHEET_STOCK_LOCATIONS = 'Warehouse Locations'
SHEET_COMPANIES = 'Companies'
SHEET_OFFERS = 'Supplier Offers'
SHEET_STOCK_ITEMS = 'Stock Items'
SHEET_HISTORY = 'Maintenance History'
SHEET_WORK_ORDERS = 'Work Orders'
# Read for their row counts only; they describe work the application cannot
# take from a spreadsheet (accounts, governed procedures, files on disk).
SHEET_STAFF = 'Staff Roles'
SHEET_PM_PLANS = 'PM Plans'
SHEET_ATTACHMENTS = 'Attachment Plan'

REQUIRED_SHEETS = (SHEET_LOCATIONS, SHEET_MACHINES)
OPTIONAL_SHEETS = (
    SHEET_PROFILES,
    SHEET_COMPONENTS,
    SHEET_PART_LINKS,
    SHEET_PARTS,
    SHEET_STOCK_LOCATIONS,
    SHEET_COMPANIES,
    SHEET_OFFERS,
    SHEET_STOCK_ITEMS,
    SHEET_HISTORY,
    SHEET_WORK_ORDERS,
)
MANUAL_SHEETS = (SHEET_STAFF, SHEET_PM_PLANS, SHEET_ATTACHMENTS)

STOCK_LOCATION_COLUMNS = ('stock_code', 'name', 'parent_stock_code', 'structural')
STOCK_LOCATION_OPTIONAL_COLUMNS = ('description',)
COMPANY_COLUMNS = ('name', 'is_manufacturer', 'is_supplier')
COMPANY_OPTIONAL_COLUMNS = (
    'company_key',
    'description',
    'website',
    'currency',
    'phone',
    'email',
    'active',
    'notes',
)
OFFER_COLUMNS = ('part_ipn', 'manufacturer', 'manufacturer_mpn')
OFFER_OPTIONAL_COLUMNS = ('supplier', 'supplier_sku', 'source_url', 'status')
STOCK_ITEM_COLUMNS = ('part_ipn', 'location_code', 'quantity_on_hand')
STOCK_ITEM_OPTIONAL_COLUMNS = (
    'stock_record_key',
    'batch',
    'status',
    'illustrative_unit_cost_usd',
    'currency',
    'financial_basis',
    'receipt_status',
    'notes',
)
HISTORY_COLUMNS = ('machine_key', 'date', 'summary')
HISTORY_OPTIONAL_COLUMNS = (
    'history_id',
    'details',
    'performed_by',
    'work_order_reference',
    'import_state',
)
WORK_ORDER_COLUMNS = (
    'work_order_key',
    'machine_key',
    'title',
    'work_order_type',
    'priority',
    'proposed_status',
)
WORK_ORDER_OPTIONAL_COLUMNS = (
    'fault_summary',
    'symptom',
    'criticality',
    'production_impact',
    'assignee',
    'due_date',
    'estimated_minutes',
    'create_repair_packet',
    'parts_list',
    'scenario_id',
    'notes',
)
MAX_STOCK_LOCATION_NAME = 100
MAX_COMPANY_NAME = 100
MAX_COMPANY_DESCRIPTION = 250
MAX_MPN = 100
MAX_BATCH = 100
MAX_HISTORY_SUMMARY = 255
MAX_HISTORY_PERFORMED_BY = 255
MAX_WORK_ORDER_TITLE = 200
MAX_WORK_ORDER_KEY = 64
MAX_WORK_ORDER_ASSIGNEE = 120
WORK_ORDER_TYPES = frozenset({
    'corrective',
    'preventive',
    'inspection',
    'calibration',
    'other',
})
WORK_ORDER_PRIORITIES = frozenset({'low', 'medium', 'high'})
# A workbook stages jobs; it does not execute them. Only states that precede
# any work are accepted, spelled either as the lifecycle value or as the
# package's "<STATE>_ENTRY_NOT_CREATED" marker.
WORK_ORDER_TARGETS = {
    'draft': 'draft',
    'planned': 'planned',
    'draft_entry_not_created': 'draft',
    'planned_entry_not_created': 'planned',
}
STOCK_STATUSES = {'ok': 10, 'attention': 50, 'damaged': 55, 'quarantined': 75}

LOCATION_COLUMNS = (
    'code',
    'name',
    'kind',
    'parent_code',
    'timezone',
    'description',
    'client_code',
    'archived',
    'reason_for_change',
)
MACHINE_COLUMNS = (
    'machine_key',
    'name',
    'asset_type',
    'asset_location_code',
    'legacy_location_label',
    'manufacturer',
    'model',
    'serial',
    'description',
    'active',
    'client_code',
    'source_key',
    'source_namespace',
    'source_entity_uuid',
    'criticality',
    'maintenance_strategy',
    'entry_status',
)
PROFILE_COLUMNS = (
    'machine_key',
    'criticality',
    'maintenance_strategy',
    'components_json',
    'energy_sources_json',
    'fault_codes_json',
    'approved_spares_json',
)
COMPONENT_COLUMNS = (
    'owner_machine_key',
    'occurrence_code',
    'display_name',
    'catalogue_part_ipn',
    'quantity',
)
PART_LINK_COLUMNS = ('machine_key', 'part_ipn', 'quantity', 'notes')
MACHINE_OPTIONAL_COLUMNS = ('public_uuid',)
COMPONENT_OPTIONAL_COLUMNS = (
    'engineering_status',
    'actual_installed_compatibility',
    'note',
)
PART_COLUMNS = (
    'ipn',
    'name',
    'category',
    'description',
    'virtual',
    'purchaseable',
    'trackable',
    'consumable',
    'units',
    'minimum_stock',
)
PART_OPTIONAL_COLUMNS = (
    'part_type',
    'manufacturer',
    'manufacturer_part_number',
    'source_url',
    'notes',
)
CATEGORY_SEPARATOR = ' / '
MAX_PART_NAME = 100
MAX_PART_DESCRIPTION = 250
MAX_PART_IPN = 100
MAX_PART_UNITS = 20
MAX_CATEGORY_NAME = 100

LOCATION_KINDS = frozenset({
    'site',
    'facility',
    'building',
    'area',
    'line',
    'cell',
    'room',
    'other',
})
ASSET_TYPES = frozenset({'equipment', 'pumphouse', 'pump'})
PUMP_KEY = re.compile(r'P[1-9][0-9]{0,3}')
SLUG = re.compile(r'^[-a-zA-Z0-9_]+$')


class WorkbookError(ValueError):
    """The workbook cannot be applied; the message names the sheet and row."""


@dataclass(frozen=True)
class LocationRecord:
    """One physical-location node to upsert."""

    row: int
    code: str
    name: str
    kind: str
    parent_code: str | None
    timezone: str
    description: str
    client_code: str
    archived: bool
    reason: str


@dataclass(frozen=True)
class MachineRecord:
    """One machine row; identity fields are only set for registry types."""

    row: int
    key: str
    name: str
    asset_type: str
    location_code: str | None
    legacy_location: str
    manufacturer: str
    model: str
    serial: str
    description: str
    active: bool
    client_code: str
    source_key: str
    source_namespace: str
    source_entity_uuid: UUID | None
    criticality: str
    maintenance_strategy: str
    entry_status: str
    public_uuid: UUID | None = None


@dataclass(frozen=True)
class ProfileRecord:
    """Knowledge profile for one machine, already JSON-decoded."""

    row: int
    machine_key: str
    profile: dict


@dataclass(frozen=True)
class ComponentRecord:
    """A component occurrence to place under a machine once its Part exists."""

    row: int
    machine_key: str
    code: str
    name: str
    part_ipn: str
    quantity: int
    review_note: str = ''


@dataclass(frozen=True)
class PartLinkRecord:
    """An installed-part quantity to record once its Part exists."""

    row: int
    machine_key: str
    part_ipn: str
    quantity: int
    notes: str


@dataclass(frozen=True)
class PartRecord:
    """A reference or catalogue Part; ``category`` is a ``A / B / C`` path."""

    row: int
    ipn: str
    name: str
    category: tuple[str, ...]
    description: str
    virtual: bool
    purchaseable: bool
    trackable: bool
    consumable: bool
    units: str
    minimum_stock: int
    part_type: str
    manufacturer: str
    mpn: str
    link: str
    notes: str


@dataclass(frozen=True)
class StockLocationRecord:
    """One warehouse node; a storage tree, never a machine's physical place."""

    row: int
    code: str
    name: str
    parent_code: str | None
    structural: bool
    description: str


@dataclass(frozen=True)
class CompanyRecord:
    """A manufacturer or supplier identity; no offer or price is implied."""

    row: int
    key: str
    name: str
    description: str
    website: str
    currency: str
    phone: str
    email: str
    is_manufacturer: bool
    is_supplier: bool
    active: bool
    notes: str


@dataclass(frozen=True)
class OfferRecord:
    """A manufacturer product identity for a catalogue part.

    A supplier offer (SKU, price, lead time) is only staged when the sheet
    names a supplier *and* a SKU; the Cedar Creek package leaves both blank
    until a distributor listing is verified, so those rows yield only the
    manufacturer part.
    """

    row: int
    part_ipn: str
    manufacturer: str
    mpn: str
    supplier: str
    sku: str
    link: str
    status: str

    @property
    def has_supplier_offer(self) -> bool:
        """True when a SupplierPart could be created from this row."""
        return bool(self.sku) and bool(self.supplier)


@dataclass(frozen=True)
class StockItemRecord:
    """Book inventory of a real catalogue part in a warehouse bin."""

    row: int
    key: str
    part_ipn: str
    location_code: str
    quantity: int
    batch: str
    status: int
    notes: str


@dataclass(frozen=True)
class HistoryRecord:
    """A dated maintenance event; never linked to a work order by import."""

    row: int
    key: str
    machine_key: str
    date: object  # datetime.date
    summary: str
    details: str
    performed_by: str


@dataclass(frozen=True)
class WorkOrderRecord:
    """A staged job in a pre-work lifecycle state."""

    row: int
    key: str
    machine_key: str
    title: str
    work_order_type: str
    priority: str
    description: str
    assignee: str
    due_date: object | None  # datetime.date
    estimated_minutes: int | None
    target_status: str
    repair_packet: bool
    scenario_id: str


@dataclass
class PopulationWorkbook:
    """Everything a population run needs, in dependency order."""

    path: str
    sha256: str
    client_codes: tuple[str, ...]
    locations: list[LocationRecord] = field(default_factory=list)
    machines: list[MachineRecord] = field(default_factory=list)
    profiles: dict[str, ProfileRecord] = field(default_factory=dict)
    components: list[ComponentRecord] = field(default_factory=list)
    part_links: list[PartLinkRecord] = field(default_factory=list)
    parts: list[PartRecord] = field(default_factory=list)
    has_parts_sheet: bool = False
    stock_locations: list[StockLocationRecord] = field(default_factory=list)
    companies: list[CompanyRecord] = field(default_factory=list)
    offers: list[OfferRecord] = field(default_factory=list)
    stock_items: list[StockItemRecord] = field(default_factory=list)
    history: list[HistoryRecord] = field(default_factory=list)
    work_orders: list[WorkOrderRecord] = field(default_factory=list)
    manual: dict[str, int] = field(default_factory=dict)

    @property
    def station(self) -> MachineRecord | None:
        """The single registered station, if the workbook has one."""
        return next((m for m in self.machines if m.asset_type == 'pumphouse'), None)

    @property
    def pumps(self) -> list[MachineRecord]:
        """Pump slots in source-key order."""
        return sorted(
            (m for m in self.machines if m.asset_type == 'pump'),
            key=lambda m: int(m.source_key[1:]),
        )

    @property
    def equipment(self) -> list[MachineRecord]:
        """Plain equipment rows in sheet order."""
        return [m for m in self.machines if m.asset_type == 'equipment']

    def summary(self) -> dict:
        """Counts the operator sees before anything is planned."""
        return {
            'workbook': self.path,
            'sha256': self.sha256,
            'client_codes': list(self.client_codes),
            'locations': len(self.locations),
            'station': self.station.name if self.station else None,
            'pumps': len(self.pumps),
            'equipment': len(self.equipment),
            'profiles': len(self.profiles),
            'parts': len(self.parts),
            'components': len(self.components),
            'part_links': len(self.part_links),
            'stock_locations': len(self.stock_locations),
            'companies': len(self.companies),
            'offers': len(self.offers),
            'stock_items': len(self.stock_items),
            'history': len(self.history),
            'work_orders': len(self.work_orders),
            'manual': dict(self.manual),
        }

    @property
    def referenced_ipns(self) -> set[str]:
        """Every IPN the component and part-link stages depend on."""
        return {c.part_ipn for c in self.components} | {
            link.part_ipn for link in self.part_links
        }


# -- cell coercion -----------------------------------------------------------


def _text(value) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _bool(value, *, sheet: str, row: int, column: str) -> bool:
    if isinstance(value, bool):
        return value
    text = _text(value).lower()
    if text in ('true', 'yes', '1', 'y'):
        return True
    if text in ('false', 'no', '0', 'n', ''):
        return False
    raise WorkbookError(
        f'{sheet} row {row}: {column} must be true/false, got {value!r}'
    )


def _int(value, *, sheet: str, row: int, column: str, minimum: int = 0) -> int:
    text = _text(value)
    if not text:
        return minimum
    try:
        number = int(float(text))
    except ValueError as exc:
        raise WorkbookError(
            f'{sheet} row {row}: {column} must be an integer, got {value!r}'
        ) from exc
    if number < minimum:
        raise WorkbookError(f'{sheet} row {row}: {column} must be >= {minimum}')
    return number


def _json_list(value, *, sheet: str, row: int, column: str) -> list:
    text = _text(value)
    if not text:
        return []
    try:
        decoded = json.loads(text)
    except json.JSONDecodeError as exc:
        raise WorkbookError(f'{sheet} row {row}: {column} is not valid JSON') from exc
    if not isinstance(decoded, list):
        raise WorkbookError(f'{sheet} row {row}: {column} must be a JSON list')
    return decoded


def _date(value, *, sheet: str, row: int, column: str, required: bool = False):
    """A calendar date from a date cell, a datetime cell or ISO text."""
    import datetime

    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = _text(value)
    if not text:
        if required:
            raise WorkbookError(f'{sheet} row {row}: {column} is required')
        return None
    try:
        return datetime.date.fromisoformat(text[:10])
    except ValueError as exc:
        raise WorkbookError(
            f'{sheet} row {row}: {column} must be a date (YYYY-MM-DD), got {value!r}'
        ) from exc


# -- sheet access ------------------------------------------------------------


def _rows(
    workbook, sheet: str, columns: tuple[str, ...], optional: tuple[str, ...] = ()
):
    """Yield (row_number, {column: value}) with the header validated.

    ``optional`` columns are read when present and yield ``None`` otherwise.
    """
    if sheet not in workbook.sheetnames:
        raise WorkbookError(f'Missing sheet: {sheet}')
    ws = workbook[sheet]
    iterator = ws.iter_rows(values_only=True)
    try:
        header = [_text(cell) for cell in next(iterator)]
    except StopIteration as exc:
        raise WorkbookError(f'{sheet}: sheet is empty') from exc
    missing = [column for column in columns if column not in header]
    if missing:
        raise WorkbookError(f'{sheet}: missing columns {", ".join(missing)}')
    index = {column: header.index(column) for column in columns}
    for column in optional:
        if column in header:
            index[column] = header.index(column)
    count = 0
    for offset, values in enumerate(iterator, start=2):
        if values is None or all(v is None or _text(v) == '' for v in values):
            continue
        count += 1
        if count > MAX_ROWS_PER_SHEET:
            raise WorkbookError(f'{sheet}: more than {MAX_ROWS_PER_SHEET} rows')
        padded = list(values) + [None] * (len(header) - len(values))
        cell = {column: padded[position] for column, position in index.items()}
        for column in optional:
            cell.setdefault(column, None)
        yield offset, cell


def _load_locations(workbook) -> list[LocationRecord]:
    sheet = SHEET_LOCATIONS
    records: list[LocationRecord] = []
    codes: dict[str, int] = {}
    for row, cell in _rows(workbook, sheet, LOCATION_COLUMNS):
        code = _text(cell['code'])
        if not code or not SLUG.match(code) or len(code) > 64:
            raise WorkbookError(f'{sheet} row {row}: code must be a slug <= 64 chars')
        key = code.lower()
        if key in codes:
            raise WorkbookError(
                f'{sheet} row {row}: code {code} duplicates row {codes[key]}'
            )
        codes[key] = row
        name = ' '.join(_text(cell['name']).split())
        if not name or len(name) > 255:
            raise WorkbookError(f'{sheet} row {row}: name is required (<= 255 chars)')
        kind = _text(cell['kind']).lower()
        if kind not in LOCATION_KINDS:
            raise WorkbookError(f'{sheet} row {row}: unknown kind {kind!r}')
        client = _text(cell['client_code'])
        if not client:
            raise WorkbookError(f'{sheet} row {row}: client_code is required')
        parent = _text(cell['parent_code']).lower() or None
        timezone = _text(cell['timezone'])
        if parent is None and not timezone:
            raise WorkbookError(f'{sheet} row {row}: a root location needs a timezone')
        description = _text(cell['description'])
        if len(description) > 10_000:
            raise WorkbookError(f'{sheet} row {row}: description exceeds 10000 chars')
        reason = _text(cell['reason_for_change']) or 'Workbook population'
        records.append(
            LocationRecord(
                row=row,
                code=key,
                name=name,
                kind=kind,
                parent_code=parent,
                timezone=timezone,
                description=description,
                client_code=client,
                archived=_bool(
                    cell['archived'], sheet=sheet, row=row, column='archived'
                ),
                reason=reason[:1000],
            )
        )
    # Parents must exist, be in the same workspace, and form a forest.
    by_code = {record.code: record for record in records}
    for record in records:
        if record.parent_code is None:
            continue
        parent = by_code.get(record.parent_code)
        if parent is None:
            raise WorkbookError(
                f'{sheet} row {record.row}: parent {record.parent_code} is not in the sheet'
            )
        if parent.client_code != record.client_code:
            raise WorkbookError(
                f'{sheet} row {record.row}: parent belongs to another client'
            )
        seen = {record.code}
        cursor = parent
        while cursor is not None:
            if cursor.code in seen:
                raise WorkbookError(
                    f'{sheet} row {record.row}: parent chain contains a cycle'
                )
            seen.add(cursor.code)
            if len(seen) > 32:
                raise WorkbookError(f'{sheet} row {record.row}: depth exceeds 32')
            cursor = by_code.get(cursor.parent_code) if cursor.parent_code else None
    return _topological(records)


def _topological(records: list[LocationRecord]) -> list[LocationRecord]:
    """Parents before children, preserving sheet order among siblings."""
    ordered: list[LocationRecord] = []
    placed: set[str] = set()
    pending = list(records)
    while pending:
        progress = False
        remaining = []
        for record in pending:
            if record.parent_code is None or record.parent_code in placed:
                ordered.append(record)
                placed.add(record.code)
                progress = True
            else:
                remaining.append(record)
        if not progress:
            raise WorkbookError(f'{SHEET_LOCATIONS}: unresolved parent chain')
        pending = remaining
    return ordered


def _load_machines(workbook, location_codes: set[str]) -> list[MachineRecord]:
    sheet = SHEET_MACHINES
    records: list[MachineRecord] = []
    keys: dict[str, int] = {}
    names: dict[str, int] = {}
    for row, cell in _rows(workbook, sheet, MACHINE_COLUMNS, MACHINE_OPTIONAL_COLUMNS):
        key = _text(cell['machine_key'])
        if not key:
            raise WorkbookError(f'{sheet} row {row}: machine_key is required')
        if key in keys:
            raise WorkbookError(
                f'{sheet} row {row}: machine_key {key} duplicates row {keys[key]}'
            )
        keys[key] = row
        name = ' '.join(_text(cell['name']).split())
        if not name or len(name) > 255:
            raise WorkbookError(f'{sheet} row {row}: name is required (<= 255 chars)')
        if name.casefold() in names:
            raise WorkbookError(
                f'{sheet} row {row}: name duplicates row {names[name.casefold()]}'
            )
        names[name.casefold()] = row
        asset_type = _text(cell['asset_type']).lower()
        if asset_type not in ASSET_TYPES:
            raise WorkbookError(f'{sheet} row {row}: unknown asset_type {asset_type!r}')
        client = _text(cell['client_code'])
        if not client:
            raise WorkbookError(f'{sheet} row {row}: client_code is required')
        location = _text(cell['asset_location_code']).lower() or None
        if location is not None and location not in location_codes:
            raise WorkbookError(
                f'{sheet} row {row}: asset_location_code {location} is not in '
                f'{SHEET_LOCATIONS}'
            )
        source_key = _text(cell['source_key'])
        namespace = _text(cell['source_namespace'])
        raw_uuid = _text(cell['source_entity_uuid'])
        entity_uuid = None
        if asset_type == 'pumphouse':
            if not source_key or not namespace or not raw_uuid:
                raise WorkbookError(
                    f'{sheet} row {row}: a station needs source_key, source_namespace '
                    'and source_entity_uuid'
                )
            if not SLUG.match(namespace) or len(namespace) > 64:
                raise WorkbookError(
                    f'{sheet} row {row}: source_namespace must be a slug'
                )
            try:
                entity_uuid = UUID(raw_uuid)
            except ValueError as exc:
                raise WorkbookError(
                    f'{sheet} row {row}: source_entity_uuid is not a UUID'
                ) from exc
        elif asset_type == 'pump':
            if not PUMP_KEY.fullmatch(source_key):
                raise WorkbookError(
                    f'{sheet} row {row}: pump source_key must be P1..P9999, got '
                    f'{source_key!r}'
                )
        for column in ('manufacturer', 'model', 'serial', 'legacy_location_label'):
            if len(_text(cell[column])) > 255:
                raise WorkbookError(f'{sheet} row {row}: {column} exceeds 255 chars')
        public_uuid = None
        raw_public = _text(cell.get('public_uuid'))
        if raw_public:
            if asset_type != 'pumphouse':
                raise WorkbookError(
                    f'{sheet} row {row}: public_uuid may only be reserved for the station'
                )
            try:
                public_uuid = UUID(raw_public)
            except ValueError as exc:
                raise WorkbookError(
                    f'{sheet} row {row}: public_uuid is not a UUID'
                ) from exc
        records.append(
            MachineRecord(
                row=row,
                key=key,
                name=name,
                asset_type=asset_type,
                location_code=location,
                legacy_location=_text(cell['legacy_location_label']),
                manufacturer=_text(cell['manufacturer']),
                model=_text(cell['model']),
                serial=_text(cell['serial']),
                description=_text(cell['description']),
                active=_bool(cell['active'], sheet=sheet, row=row, column='active'),
                client_code=client,
                source_key=source_key,
                source_namespace=namespace,
                source_entity_uuid=entity_uuid,
                criticality=_text(cell['criticality']).lower(),
                maintenance_strategy=_text(cell['maintenance_strategy']).lower(),
                entry_status=_text(cell['entry_status']),
                public_uuid=public_uuid,
            )
        )
    stations = [m for m in records if m.asset_type == 'pumphouse']
    pumps = [m for m in records if m.asset_type == 'pump']
    if len(stations) > 1:
        raise WorkbookError(
            f'{sheet}: more than one pumphouse row; one station per run'
        )
    if pumps and not stations:
        raise WorkbookError(
            f'{sheet}: pump rows need a pumphouse row in the same sheet'
        )
    pump_keys: dict[str, int] = {}
    for pump in pumps:
        if pump.source_key in pump_keys:
            raise WorkbookError(
                f'{sheet} row {pump.row}: pump source_key {pump.source_key} duplicates '
                f'row {pump_keys[pump.source_key]}'
            )
        pump_keys[pump.source_key] = pump.row
        if pump.client_code != stations[0].client_code:
            raise WorkbookError(
                f'{sheet} row {pump.row}: pump client differs from its station'
            )
    return records


def _load_profiles(workbook, machine_keys: set[str]) -> dict[str, ProfileRecord]:
    sheet = SHEET_PROFILES
    if sheet not in workbook.sheetnames:
        return {}
    profiles: dict[str, ProfileRecord] = {}
    for row, cell in _rows(workbook, sheet, PROFILE_COLUMNS):
        key = _text(cell['machine_key'])
        if key not in machine_keys:
            raise WorkbookError(
                f'{sheet} row {row}: machine_key {key} is not in {SHEET_MACHINES}'
            )
        if key in profiles:
            raise WorkbookError(
                f'{sheet} row {row}: machine_key {key} duplicates row '
                f'{profiles[key].row}'
            )
        profile: dict = {}
        criticality = _text(cell['criticality']).lower()
        strategy = _text(cell['maintenance_strategy']).lower()
        if criticality:
            profile['criticality'] = criticality
        if strategy:
            profile['maintenance_strategy'] = strategy
        for column, field_name in (
            ('components_json', 'components'),
            ('energy_sources_json', 'energy_sources'),
            ('fault_codes_json', 'fault_codes'),
            ('approved_spares_json', 'approved_spares'),
        ):
            profile[field_name] = _json_list(
                cell[column], sheet=sheet, row=row, column=column
            )
        profiles[key] = ProfileRecord(row=row, machine_key=key, profile=profile)
    return profiles


def _load_components(workbook, machine_keys: set[str]) -> list[ComponentRecord]:
    sheet = SHEET_COMPONENTS
    if sheet not in workbook.sheetnames:
        return []
    records: list[ComponentRecord] = []
    slots: dict[tuple[str, str], int] = {}
    for row, cell in _rows(
        workbook, sheet, COMPONENT_COLUMNS, COMPONENT_OPTIONAL_COLUMNS
    ):
        key = _text(cell['owner_machine_key'])
        if key not in machine_keys:
            raise WorkbookError(
                f'{sheet} row {row}: owner_machine_key {key} is not in {SHEET_MACHINES}'
            )
        code = _text(cell['occurrence_code'])
        if not code or len(code) > 100:
            raise WorkbookError(f'{sheet} row {row}: occurrence_code required (<= 100)')
        if (key, code) in slots:
            raise WorkbookError(
                f'{sheet} row {row}: occurrence {code} on {key} duplicates row '
                f'{slots[key, code]}'
            )
        slots[key, code] = row
        ipn = _text(cell['catalogue_part_ipn'])
        if not ipn:
            raise WorkbookError(f'{sheet} row {row}: catalogue_part_ipn is required')
        note_parts = [
            _text(cell.get('note')),
            ' '.join(
                f'{label}={_text(cell.get(column))}'
                for label, column in (
                    ('engineering', 'engineering_status'),
                    ('compatibility', 'actual_installed_compatibility'),
                )
                if _text(cell.get(column))
            ),
        ]
        records.append(
            ComponentRecord(
                row=row,
                machine_key=key,
                code=code,
                name=_text(cell['display_name'])[:255] or code,
                part_ipn=ipn,
                quantity=_int(
                    cell['quantity'], sheet=sheet, row=row, column='quantity', minimum=1
                ),
                review_note=' '.join(p for p in note_parts if p)[:4000],
            )
        )
    return records


def _load_part_links(workbook, machine_keys: set[str]) -> list[PartLinkRecord]:
    sheet = SHEET_PART_LINKS
    if sheet not in workbook.sheetnames:
        return []
    records: list[PartLinkRecord] = []
    pairs: dict[tuple[str, str], int] = {}
    for row, cell in _rows(workbook, sheet, PART_LINK_COLUMNS):
        key = _text(cell['machine_key'])
        if key not in machine_keys:
            raise WorkbookError(
                f'{sheet} row {row}: machine_key {key} is not in {SHEET_MACHINES}'
            )
        ipn = _text(cell['part_ipn'])
        if not ipn:
            raise WorkbookError(f'{sheet} row {row}: part_ipn is required')
        if (key, ipn) in pairs:
            raise WorkbookError(
                f'{sheet} row {row}: ({key}, {ipn}) duplicates row {pairs[key, ipn]}'
            )
        pairs[key, ipn] = row
        records.append(
            PartLinkRecord(
                row=row,
                machine_key=key,
                part_ipn=ipn,
                quantity=_int(
                    cell['quantity'], sheet=sheet, row=row, column='quantity', minimum=1
                ),
                notes=_text(cell['notes']),
            )
        )
    return records


def _load_parts(workbook) -> list[PartRecord]:
    sheet = SHEET_PARTS
    if sheet not in workbook.sheetnames:
        return []
    records: list[PartRecord] = []
    ipns: dict[str, int] = {}
    for row, cell in _rows(workbook, sheet, PART_COLUMNS, PART_OPTIONAL_COLUMNS):
        ipn = _text(cell['ipn'])
        if not ipn or len(ipn) > MAX_PART_IPN:
            raise WorkbookError(
                f'{sheet} row {row}: ipn is required (<= {MAX_PART_IPN})'
            )
        if ipn.casefold() in ipns:
            raise WorkbookError(
                f'{sheet} row {row}: ipn {ipn} duplicates row {ipns[ipn.casefold()]}'
            )
        ipns[ipn.casefold()] = row
        name = ' '.join(_text(cell['name']).split())
        if not name or len(name) > MAX_PART_NAME:
            raise WorkbookError(
                f'{sheet} row {row}: name is required (<= {MAX_PART_NAME} chars)'
            )
        description = _text(cell['description'])
        if len(description) > MAX_PART_DESCRIPTION:
            raise WorkbookError(
                f'{sheet} row {row}: description exceeds {MAX_PART_DESCRIPTION} chars'
            )
        raw_category = _text(cell['category'])
        category = tuple(
            ' '.join(segment.split())
            for segment in raw_category.split(CATEGORY_SEPARATOR)
            if segment.strip()
        )
        if not category:
            raise WorkbookError(f'{sheet} row {row}: category is required')
        if any(len(segment) > MAX_CATEGORY_NAME for segment in category):
            raise WorkbookError(
                f'{sheet} row {row}: category segment exceeds {MAX_CATEGORY_NAME} chars'
            )
        units = _text(cell['units'])
        if len(units) > MAX_PART_UNITS:
            raise WorkbookError(
                f'{sheet} row {row}: units exceeds {MAX_PART_UNITS} chars'
            )
        virtual = _bool(cell['virtual'], sheet=sheet, row=row, column='virtual')
        purchaseable = _bool(
            cell['purchaseable'], sheet=sheet, row=row, column='purchaseable'
        )
        if virtual and purchaseable:
            raise WorkbookError(
                f'{sheet} row {row}: a virtual reference part cannot be purchaseable'
            )
        records.append(
            PartRecord(
                row=row,
                ipn=ipn,
                name=name,
                category=category,
                description=description,
                virtual=virtual,
                purchaseable=purchaseable,
                trackable=_bool(
                    cell['trackable'], sheet=sheet, row=row, column='trackable'
                ),
                consumable=_bool(
                    cell['consumable'], sheet=sheet, row=row, column='consumable'
                ),
                units=units,
                minimum_stock=_int(
                    cell['minimum_stock'], sheet=sheet, row=row, column='minimum_stock'
                ),
                part_type=_text(cell.get('part_type')),
                manufacturer=_text(cell.get('manufacturer'))[:250],
                mpn=_text(cell.get('manufacturer_part_number'))[:100],
                link=_text(cell.get('source_url'))[:2000],
                notes=_text(cell.get('notes')),
            )
        )
    return records


def _load_stock_locations(workbook) -> list[StockLocationRecord]:
    sheet = SHEET_STOCK_LOCATIONS
    if sheet not in workbook.sheetnames:
        return []
    records: list[StockLocationRecord] = []
    codes: dict[str, int] = {}
    for row, cell in _rows(
        workbook, sheet, STOCK_LOCATION_COLUMNS, STOCK_LOCATION_OPTIONAL_COLUMNS
    ):
        code = _text(cell['stock_code']).lower()
        if not code:
            raise WorkbookError(f'{sheet} row {row}: stock_code is required')
        if code in codes:
            raise WorkbookError(
                f'{sheet} row {row}: stock_code {code} duplicates row {codes[code]}'
            )
        codes[code] = row
        name = ' '.join(_text(cell['name']).split())
        if not name or len(name) > MAX_STOCK_LOCATION_NAME:
            raise WorkbookError(
                f'{sheet} row {row}: name is required (<= {MAX_STOCK_LOCATION_NAME})'
            )
        parent = _text(cell['parent_stock_code']).lower() or None
        if parent == code:
            raise WorkbookError(f'{sheet} row {row}: a location cannot parent itself')
        records.append(
            StockLocationRecord(
                row=row,
                code=code,
                name=name,
                parent_code=parent,
                structural=_bool(
                    cell['structural'], sheet=sheet, row=row, column='structural'
                ),
                description=_text(cell.get('description'))[:250],
            )
        )
    # Parents must be defined, before their children, and acyclic.
    seen: set[str] = set()
    for record in records:
        if record.parent_code is not None and record.parent_code not in seen:
            if record.parent_code in codes:
                raise WorkbookError(
                    f'{sheet} row {record.row}: parent {record.parent_code} must '
                    'appear before its children'
                )
            raise WorkbookError(
                f'{sheet} row {record.row}: parent {record.parent_code} is not in the sheet'
            )
        seen.add(record.code)
    return records


def _load_companies(workbook) -> list[CompanyRecord]:
    sheet = SHEET_COMPANIES
    if sheet not in workbook.sheetnames:
        return []
    records: list[CompanyRecord] = []
    names: dict[str, int] = {}
    for row, cell in _rows(workbook, sheet, COMPANY_COLUMNS, COMPANY_OPTIONAL_COLUMNS):
        name = ' '.join(_text(cell['name']).split())
        if not name or len(name) > MAX_COMPANY_NAME:
            raise WorkbookError(
                f'{sheet} row {row}: name is required (<= {MAX_COMPANY_NAME})'
            )
        if name.casefold() in names:
            raise WorkbookError(
                f'{sheet} row {row}: name {name!r} duplicates row {names[name.casefold()]}'
            )
        names[name.casefold()] = row
        description = _text(cell.get('description'))
        if len(description) > MAX_COMPANY_DESCRIPTION:
            raise WorkbookError(
                f'{sheet} row {row}: description exceeds {MAX_COMPANY_DESCRIPTION} chars'
            )
        currency = _text(cell.get('currency')).upper()
        if currency and len(currency) != 3:
            raise WorkbookError(f'{sheet} row {row}: currency must be a 3-letter code')
        is_manufacturer = _bool(
            cell['is_manufacturer'], sheet=sheet, row=row, column='is_manufacturer'
        )
        is_supplier = _bool(
            cell['is_supplier'], sheet=sheet, row=row, column='is_supplier'
        )
        if not is_manufacturer and not is_supplier:
            raise WorkbookError(
                f'{sheet} row {row}: a company must be a manufacturer or a supplier'
            )
        active_cell = cell.get('active')
        records.append(
            CompanyRecord(
                row=row,
                key=_text(cell.get('company_key')),
                name=name,
                description=description,
                website=_text(cell.get('website'))[:200],
                currency=currency,
                phone=_text(cell.get('phone'))[:50],
                email=_text(cell.get('email'))[:254],
                is_manufacturer=is_manufacturer,
                is_supplier=is_supplier,
                active=(
                    True
                    if active_cell is None
                    else _bool(active_cell, sheet=sheet, row=row, column='active')
                ),
                notes=_text(cell.get('notes')),
            )
        )
    return records


def _load_offers(workbook) -> list[OfferRecord]:
    sheet = SHEET_OFFERS
    if sheet not in workbook.sheetnames:
        return []
    records: list[OfferRecord] = []
    seen: dict[tuple[str, str, str], int] = {}
    for row, cell in _rows(workbook, sheet, OFFER_COLUMNS, OFFER_OPTIONAL_COLUMNS):
        ipn = _text(cell['part_ipn'])
        manufacturer = ' '.join(_text(cell['manufacturer']).split())
        mpn = _text(cell['manufacturer_mpn'])
        if not ipn or not manufacturer or not mpn:
            raise WorkbookError(
                f'{sheet} row {row}: part_ipn, manufacturer and manufacturer_mpn are required'
            )
        if len(mpn) > MAX_MPN:
            raise WorkbookError(
                f'{sheet} row {row}: manufacturer_mpn exceeds {MAX_MPN}'
            )
        identity = (ipn.casefold(), manufacturer.casefold(), mpn.casefold())
        if identity in seen:
            raise WorkbookError(
                f'{sheet} row {row}: offer duplicates row {seen[identity]}'
            )
        seen[identity] = row
        supplier = ' '.join(_text(cell.get('supplier')).split())
        if supplier.upper() in ('', 'NOT_SELECTED', 'NONE', 'TBD'):
            supplier = ''
        records.append(
            OfferRecord(
                row=row,
                part_ipn=ipn,
                manufacturer=manufacturer,
                mpn=mpn,
                supplier=supplier,
                sku=_text(cell.get('supplier_sku'))[:100],
                link=_text(cell.get('source_url'))[:2000],
                status=_text(cell.get('status')),
            )
        )
    return records


def _load_stock_items(workbook) -> list[StockItemRecord]:
    sheet = SHEET_STOCK_ITEMS
    if sheet not in workbook.sheetnames:
        return []
    records: list[StockItemRecord] = []
    seen: dict[tuple[str, str, str], int] = {}
    for row, cell in _rows(
        workbook, sheet, STOCK_ITEM_COLUMNS, STOCK_ITEM_OPTIONAL_COLUMNS
    ):
        ipn = _text(cell['part_ipn'])
        location = _text(cell['location_code']).lower()
        if not ipn or not location:
            raise WorkbookError(
                f'{sheet} row {row}: part_ipn and location_code are required'
            )
        batch = _text(cell.get('batch'))[:MAX_BATCH]
        identity = (ipn.casefold(), location, batch.casefold())
        if identity in seen:
            raise WorkbookError(
                f'{sheet} row {row}: (part, location, batch) duplicates row {seen[identity]}'
            )
        seen[identity] = row
        status_text = _text(cell.get('status')).lower() or 'ok'
        if status_text not in STOCK_STATUSES:
            raise WorkbookError(
                f'{sheet} row {row}: status must be one of '
                + ', '.join(sorted(STOCK_STATUSES))
            )
        # The illustrative cost is kept as text in the notes, never as a
        # purchase price: the package marks it as an estimate, not an offer.
        notes = []
        cost = _text(cell.get('illustrative_unit_cost_usd'))
        if cost:
            notes.append(
                f'Illustrative unit cost {cost} {_text(cell.get("currency")) or "USD"}'
            )
        for column in ('financial_basis', 'receipt_status'):
            value = _text(cell.get(column))
            if value:
                notes.append(f'{column}={value}')
        free = _text(cell.get('notes'))
        if free:
            notes.append(free)
        records.append(
            StockItemRecord(
                row=row,
                key=_text(cell.get('stock_record_key')),
                part_ipn=ipn,
                location_code=location,
                quantity=_int(
                    cell['quantity_on_hand'],
                    sheet=sheet,
                    row=row,
                    column='quantity_on_hand',
                ),
                batch=batch,
                status=STOCK_STATUSES[status_text],
                notes='; '.join(notes),
            )
        )
    return records


def _load_history(workbook, machine_keys: set[str]) -> list[HistoryRecord]:
    sheet = SHEET_HISTORY
    if sheet not in workbook.sheetnames:
        return []
    records: list[HistoryRecord] = []
    seen: dict[tuple[str, object, str], int] = {}
    for row, cell in _rows(workbook, sheet, HISTORY_COLUMNS, HISTORY_OPTIONAL_COLUMNS):
        key = _text(cell['machine_key'])
        if key not in machine_keys:
            raise WorkbookError(
                f'{sheet} row {row}: machine_key {key} is not in {SHEET_MACHINES}'
            )
        if _text(cell.get('work_order_reference')):
            raise WorkbookError(
                f'{sheet} row {row}: history rows cannot reference a work order; '
                'a closed job is recorded by closing it, not by import'
            )
        summary = ' '.join(_text(cell['summary']).split())
        if not summary or len(summary) > MAX_HISTORY_SUMMARY:
            raise WorkbookError(
                f'{sheet} row {row}: summary is required (<= {MAX_HISTORY_SUMMARY})'
            )
        date = _date(cell['date'], sheet=sheet, row=row, column='date', required=True)
        identity = (key, date, summary.casefold())
        if identity in seen:
            raise WorkbookError(
                f'{sheet} row {row}: (machine, date, summary) duplicates row {seen[identity]}'
            )
        seen[identity] = row
        records.append(
            HistoryRecord(
                row=row,
                key=_text(cell.get('history_id')),
                machine_key=key,
                date=date,
                summary=summary,
                details=_text(cell.get('details')),
                performed_by=_text(cell.get('performed_by'))[:MAX_HISTORY_PERFORMED_BY],
            )
        )
    return records


def _load_work_orders(workbook, machine_keys: set[str]) -> list[WorkOrderRecord]:
    sheet = SHEET_WORK_ORDERS
    if sheet not in workbook.sheetnames:
        return []
    records: list[WorkOrderRecord] = []
    keys: dict[str, int] = {}
    for row, cell in _rows(
        workbook, sheet, WORK_ORDER_COLUMNS, WORK_ORDER_OPTIONAL_COLUMNS
    ):
        key = _text(cell['work_order_key'])
        if not key or len(key) > MAX_WORK_ORDER_KEY:
            raise WorkbookError(
                f'{sheet} row {row}: work_order_key is required (<= {MAX_WORK_ORDER_KEY})'
            )
        if key in keys:
            raise WorkbookError(
                f'{sheet} row {row}: work_order_key {key} duplicates row {keys[key]}'
            )
        keys[key] = row
        machine_key = _text(cell['machine_key'])
        if machine_key not in machine_keys:
            raise WorkbookError(
                f'{sheet} row {row}: machine_key {machine_key} is not in {SHEET_MACHINES}'
            )
        title = ' '.join(_text(cell['title']).split())
        if not title or len(title) > MAX_WORK_ORDER_TITLE:
            raise WorkbookError(
                f'{sheet} row {row}: title is required (<= {MAX_WORK_ORDER_TITLE})'
            )
        kind = _text(cell['work_order_type']).lower()
        if kind not in WORK_ORDER_TYPES:
            raise WorkbookError(
                f'{sheet} row {row}: work_order_type must be one of '
                + ', '.join(sorted(WORK_ORDER_TYPES))
            )
        priority = _text(cell['priority']).lower()
        if priority not in WORK_ORDER_PRIORITIES:
            raise WorkbookError(
                f'{sheet} row {row}: priority must be one of '
                + ', '.join(sorted(WORK_ORDER_PRIORITIES))
            )
        proposed = _text(cell['proposed_status']).lower()
        if proposed not in WORK_ORDER_TARGETS:
            raise WorkbookError(
                f'{sheet} row {row}: proposed_status {proposed!r} is not a staging '
                'state (draft or planned)'
            )
        lines = []
        for column, label in (
            ('fault_summary', 'Fault'),
            ('symptom', 'Symptom'),
            ('criticality', 'Criticality'),
            ('production_impact', 'Production impact'),
            ('parts_list', 'Parts'),
            ('notes', 'Notes'),
        ):
            value = _text(cell.get(column))
            if value:
                lines.append(f'{label}: {value}')
        minutes_cell = cell.get('estimated_minutes')
        records.append(
            WorkOrderRecord(
                row=row,
                key=key,
                machine_key=machine_key,
                title=title,
                work_order_type=kind,
                priority=priority,
                description='\n'.join(lines),
                assignee=_text(cell.get('assignee'))[:MAX_WORK_ORDER_ASSIGNEE],
                due_date=_date(
                    cell.get('due_date'), sheet=sheet, row=row, column='due_date'
                ),
                estimated_minutes=(
                    None
                    if _text(minutes_cell) == ''
                    else _int(
                        minutes_cell, sheet=sheet, row=row, column='estimated_minutes'
                    )
                ),
                target_status=WORK_ORDER_TARGETS[proposed],
                repair_packet=_bool(
                    cell.get('create_repair_packet'),
                    sheet=sheet,
                    row=row,
                    column='create_repair_packet',
                ),
                scenario_id=_text(cell.get('scenario_id'))[:50],
            )
        )
    return records


def _count_rows(workbook, sheet: str) -> int:
    """Non-empty data rows of a sheet the run reports but never applies."""
    if sheet not in workbook.sheetnames:
        return 0
    rows = workbook[sheet].iter_rows(values_only=True)
    next(rows, None)
    return sum(
        1
        for values in rows
        if values is not None and any(v is not None and _text(v) != '' for v in values)
    )


def load_population_workbook(path) -> PopulationWorkbook:
    """Read and validate a population workbook from disk.

    Raises ``WorkbookError`` with a sheet/row reference for anything that
    would make an apply ambiguous. ``openpyxl`` is imported lazily so the
    module can be imported where only the dataclasses are needed.
    """
    from openpyxl import load_workbook

    file = Path(path)
    if not file.is_file():
        raise WorkbookError(f'Workbook not found: {file}')
    size = file.stat().st_size
    if size > MAX_WORKBOOK_BYTES:
        raise WorkbookError(f'Workbook exceeds {MAX_WORKBOOK_BYTES} bytes')
    digest = hashlib.sha256(file.read_bytes()).hexdigest()
    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types
        raise WorkbookError(f'Cannot open workbook: {exc}') from exc
    try:
        for sheet in REQUIRED_SHEETS:
            if sheet not in workbook.sheetnames:
                raise WorkbookError(f'Missing sheet: {sheet}')
        locations = _load_locations(workbook)
        location_codes = {record.code for record in locations}
        machines = _load_machines(workbook, location_codes)
        machine_keys = {record.key for record in machines}
        profiles = _load_profiles(workbook, machine_keys)
        components = _load_components(workbook, machine_keys)
        part_links = _load_part_links(workbook, machine_keys)
        has_parts_sheet = SHEET_PARTS in workbook.sheetnames
        parts = _load_parts(workbook)
        stock_locations = _load_stock_locations(workbook)
        companies = _load_companies(workbook)
        offers = _load_offers(workbook)
        stock_items = _load_stock_items(workbook)
        history = _load_history(workbook, machine_keys)
        work_orders = _load_work_orders(workbook, machine_keys)
        manual = {
            sheet: _count_rows(workbook, sheet)
            for sheet in MANUAL_SHEETS
            if sheet in workbook.sheetnames
        }
    finally:
        workbook.close()
    if has_parts_sheet:
        # With a catalogue present every referenced IPN must be defined by it;
        # silently falling through to "whatever is in the database" would let a
        # typo bind a component to an unrelated part.
        catalogue = {record.ipn.casefold() for record in parts}
        referenced = {c.part_ipn for c in components} | {
            link.part_ipn for link in part_links
        }
        missing = sorted(ipn for ipn in referenced if ipn.casefold() not in catalogue)
        if missing:
            raise WorkbookError(
                f'{SHEET_PARTS}: referenced IPNs are not in the catalogue: '
                + ', '.join(missing[:20])
            )
    _check_operations(
        locations=locations,
        parts=parts,
        has_parts_sheet=has_parts_sheet,
        stock_locations=stock_locations,
        companies=companies,
        offers=offers,
        stock_items=stock_items,
    )
    clients = sorted(
        {record.client_code for record in locations}
        | {record.client_code for record in machines}
    )
    if len(clients) != 1:
        raise WorkbookError(
            'A population workbook must name exactly one client_code; found '
            + (', '.join(clients) or 'none')
        )
    return PopulationWorkbook(
        path=str(file),
        sha256=digest,
        client_codes=tuple(clients),
        locations=locations,
        machines=machines,
        profiles=profiles,
        components=components,
        part_links=part_links,
        parts=parts,
        has_parts_sheet=has_parts_sheet,
        stock_locations=stock_locations,
        companies=companies,
        offers=offers,
        stock_items=stock_items,
        history=history,
        work_orders=work_orders,
        manual=manual,
    )


def _check_operations(
    *,
    locations,
    parts,
    has_parts_sheet,
    stock_locations,
    companies,
    offers,
    stock_items,
):
    """Cross-sheet rules for the inventory stages.

    * Warehouse codes never collide with physical-location codes (QA-05).
    * Stock is only booked against a real, non-virtual catalogue part in a
      non-structural bin (QA-18, QA-22).
    * A manufacturer named on an offer is a manufacturer in ``Companies``.
    """
    physical = {record.code for record in locations}
    clash = sorted(r.code for r in stock_locations if r.code in physical)
    if clash:
        raise WorkbookError(
            f'{SHEET_STOCK_LOCATIONS}: codes also used by {SHEET_LOCATIONS}: '
            + ', '.join(clash[:10])
        )
    bins = {record.code: record for record in stock_locations}
    catalogue = {record.ipn.casefold(): record for record in parts}
    for item in stock_items:
        location = bins.get(item.location_code)
        if location is None:
            raise WorkbookError(
                f'{SHEET_STOCK_ITEMS} row {item.row}: location_code '
                f'{item.location_code} is not in {SHEET_STOCK_LOCATIONS}'
            )
        if location.structural:
            raise WorkbookError(
                f'{SHEET_STOCK_ITEMS} row {item.row}: {item.location_code} is '
                'structural and cannot hold stock'
            )
        if has_parts_sheet:
            part = catalogue.get(item.part_ipn.casefold())
            if part is None:
                raise WorkbookError(
                    f'{SHEET_STOCK_ITEMS} row {item.row}: part_ipn {item.part_ipn} '
                    f'is not in {SHEET_PARTS}'
                )
            if part.virtual:
                raise WorkbookError(
                    f'{SHEET_STOCK_ITEMS} row {item.row}: {item.part_ipn} is a '
                    'virtual reference part and cannot carry stock'
                )
    manufacturers = {
        record.name.casefold() for record in companies if record.is_manufacturer
    }
    suppliers = {record.name.casefold() for record in companies if record.is_supplier}
    for offer in offers:
        if has_parts_sheet and offer.part_ipn.casefold() not in catalogue:
            raise WorkbookError(
                f'{SHEET_OFFERS} row {offer.row}: part_ipn {offer.part_ipn} is not '
                f'in {SHEET_PARTS}'
            )
        if companies and offer.manufacturer.casefold() not in manufacturers:
            raise WorkbookError(
                f'{SHEET_OFFERS} row {offer.row}: manufacturer {offer.manufacturer!r} '
                f'is not a manufacturer in {SHEET_COMPANIES}'
            )
        if offer.supplier and companies and offer.supplier.casefold() not in suppliers:
            raise WorkbookError(
                f'{SHEET_OFFERS} row {offer.row}: supplier {offer.supplier!r} is not '
                f'a supplier in {SHEET_COMPANIES}'
            )


__all__ = [
    'CompanyRecord',
    'ComponentRecord',
    'HistoryRecord',
    'LocationRecord',
    'MachineRecord',
    'OfferRecord',
    'PartLinkRecord',
    'PartRecord',
    'PopulationWorkbook',
    'ProfileRecord',
    'StockItemRecord',
    'StockLocationRecord',
    'WorkOrderRecord',
    'WorkbookError',
    'load_population_workbook',
]
